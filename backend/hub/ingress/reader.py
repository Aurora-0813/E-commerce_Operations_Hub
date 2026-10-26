# -*- coding: utf-8 -*-
"""
xlsx 读取（reader.py）
=====================
负责"把 xlsx 读成行 + 找到表头 + 识别空行/合计行"。

真实导出文件的几个坑（都在这里处理）：
    1. 表头不一定在第一行（可能有说明行）→ 扫描前 N 行找
    2. 中间偶发空行 → 跳过
    3. 末尾有"合计"行 → 跳过
"""
from pathlib import Path

from openpyxl import load_workbook


def read_sheets(path: Path) -> list[tuple[str, list[list]]]:
    """读 xlsx 的所有 sheet，返回 [(sheet 名, 行列表)]。

    每一行是一个 list（单元格值的列表）。
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    result: list[tuple[str, list[list]]] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        result.append((ws.title, rows))
    wb.close()
    return result


def detect_header(rows: list[list], known: set[str], max_scan: int = 10) -> tuple[int, dict[str, int]]:
    """探测表头行。

    做法：看前 max_scan 行，哪一行命中的"已知列名"最多，哪一行就是表头。

    :return: (表头行下标, {列名: 列下标})
             表头行下标是 0-based（方便取数据行），对外展示再 +1
    """
    best_idx, best_map, best_hits = 0, {}, 0
    for i, row in enumerate(rows[:max_scan]):
        col_map: dict[str, int] = {}
        for j, cell in enumerate(row):
            if cell is None:
                continue
            name = str(cell).strip()
            if name in known:
                col_map[name] = j
        if len(col_map) > best_hits:
            best_hits = len(col_map)
            best_idx = i
            best_map = col_map
    return best_idx, best_map


def is_empty_row(row: list) -> bool:
    """整行都是空（空行）。"""
    return all(c is None or str(c).strip() == "" for c in row)


def is_total_row(row: list) -> bool:
    """合计行 / 汇总行（导出文件末尾常见）。"""
    for c in row:
        if c is None:
            continue
        s = str(c).strip()
        if s in ("合计", "总计", "小计", "汇总") or s.startswith("合计"):
            return True
    return False
