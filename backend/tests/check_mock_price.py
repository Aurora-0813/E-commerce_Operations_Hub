# -*- coding: utf-8 -*-
"""
检查 mock 数据里的"价格背离"和"问题商品"是否造对了（不用连数据库）。

验证三件事：
    ① 5 个商品的两平台价差 |dev| > 0.15
    ② 3 个问题商品的差评率明显高于其他
    ③ 同时满足"价差大 + 差评率高"的商品 ≈ 3 个 ← 这就是 M4 该检出的数量

运行（在 backend/ 目录下，先跑 generate_mock.py）：
    python tests/check_mock_price.py
"""
import math
import sys
from pathlib import Path

_BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(_BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(_BACKEND_DIR))

import openpyxl   # noqa: E402

DATA = _BACKEND_DIR.parent / "data" / "mock"


def _price(v) -> float | None:
    """'￥199.00' → 199.0"""
    if v is None:
        return None
    s = str(v).replace("￥", "").replace("¥", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None


def load_products(path: Path, title_col: str, price_col: str) -> dict[str, float]:
    """读商品导出：{标题: 价格}"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    ti, pi = header.index(title_col), header.index(price_col)
    out = {}
    for r in rows[1:]:
        p = _price(r[pi])
        if p is not None:
            out[str(r[ti])] = p
    wb.close()
    return out


def load_review_stats(path: Path) -> dict[str, tuple[int, int]]:
    """读评价导出：{标题: (评价数, 差评数)}"""
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header = list(rows[0])
    ti, gi = header.index("商品标题"), header.index("评价等级")
    out: dict[str, list[int]] = {}
    for r in rows[1:]:
        title = str(r[ti])
        cell = out.setdefault(title, [0, 0])
        cell[0] += 1
        if str(r[gi]) == "差评":
            cell[1] += 1
    wb.close()
    return {k: (v[0], v[1]) for k, v in out.items()}


def main() -> None:
    tb = load_products(DATA / "taobao_product_202610.xlsx", "商品标题", "价格")
    jd = load_products(DATA / "jd_product_202610.xlsx", "商品名称", "京东价")
    common = sorted(set(tb) & set(jd))
    print(f"淘宝商品 {len(tb)}｜京东商品 {len(jd)}｜跨平台同款 {len(common)}\n")

    tb_stats = load_review_stats(DATA / "taobao_review_202610.xlsx")

    # 差评率的第 70 百分位（只统计评价数 >= 50 的商品，避免小样本噪声）
    rates = sorted(
        (b / n for title, (n, b) in tb_stats.items() if n >= 50), reverse=True
    )
    cut = rates[max(0, int(len(rates) * 0.30) - 1)] if rates else 0.0
    print(f"差评率第 70 百分位 ≈ {cut:.1%}（超过它 = 高差评）\n")

    rows = []
    for title in common:
        p1, p2 = tb[title], jd[title]
        dev = math.log(p2 / ((p1 + p2) / 2))       # 京东相对两平台中位的偏离
        n, b = tb_stats.get(title, (0, 0))
        rows.append((title, p1, p2, dev, b / n if n else 0.0, n))

    print(f"{'商品':<22}{'淘宝':>6}{'京东':>6}{'dev':>8}{'差评率':>8}{'判定':>10}")
    print("-" * 62)
    n_big = n_hit = 0
    for title, p1, p2, dev, rate, n in sorted(rows, key=lambda x: -abs(x[3])):
        big = abs(dev) >= 0.15
        high = rate >= cut
        tags = []
        if big:
            n_big += 1
            tags.append("价差大")
        if high:
            tags.append("高差评")
        if big and high:
            n_hit += 1
            tags.append("★命中")
        print(f"{title[:20]:<22}{p1:>6.0f}{p2:>6.0f}{dev:>8.3f}{rate:>8.1%}"
              f"{'/'.join(tags):>10}")

    print(f"\n价差 > 0.15：{n_big} 个（预期 5）")
    print(f"价差大 + 高差评：{n_hit} 个（预期 3 ← M4 应该检出这么多）")


if __name__ == "__main__":
    main()
